import io
import re
from html import escape, unescape
from urllib.parse import quote

from fastapi.responses import StreamingResponse

from backend.schemas.satellite import SSatelliteOut


def strip_tags(value: str | None) -> str | None:
    if not value:
        return value
    without_tags = re.sub(r'<[^>]+>', '', value)
    return unescape(without_tags).strip()


def content_disposition(filename: str) -> str:
    return f"attachment; filename*=UTF-8''{quote(filename)}"


def get_rows(sat: SSatelliteOut) -> list[tuple[str, str]]:
    return [
        ("Страна", ", ".join(sat.country) or "—"),
        ("Масса", f"{sat.mass} кг" if sat.mass else "—"),
        ("Диапазон", sat.frequency_range or "—"),
        ("Разрешение", sat.resolution or "—"),
        ("Радиометрическая чувствительность", sat.radiometric_sensitivity or "—"),
        ("Описание", sat.small_content or "—"),
        ("Подробнее", sat.big_content or "—"),
    ]


def export_excel(sat: SSatelliteOut) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Карточка объекта"

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="F5F5F5")

    ws.append(["Параметр", "Значение"])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    ws.append(["Название", sat.title_content or ""])
    for key, value in get_rows(sat):
        ws.append([key, value])

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 60
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"{sat.title_content or 'object'}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(filename)},
    )


def export_doc(sat: SSatelliteOut) -> StreamingResponse:
    from docx import Document
    from docx.shared import RGBColor

    doc = Document()
    doc.add_heading(sat.title_content or "", level=1)
    if sat.small_content:
        p = doc.add_paragraph(sat.small_content)
        p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Параметр"
    hdr[1].text = "Значение"
    for cell in hdr:
        cell.paragraphs[0].add_run().bold = True

    for key, value in [("Название", sat.title_content or "")] + get_rows(sat):
        row = table.add_row().cells
        row[0].text = key
        row[1].text = value

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    filename = f"{sat.title_content or 'object'}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": content_disposition(filename)},
    )


def export_pdf(sat: SSatelliteOut) -> StreamingResponse:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import os

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=20*mm, rightMargin=20*mm,
                            topMargin=20*mm, bottomMargin=20*mm)

    font_path = "dist/fonts/GolosText-Regular.ttf"
    font_bold_path = "dist/fonts/GolosText-Bold.ttf"
    if os.path.exists(font_path):
        pdfmetrics.registerFont(TTFont("Golos", font_path))
        pdfmetrics.registerFont(TTFont("GolosBold", font_bold_path))
        base_font, bold_font = "Golos", "GolosBold"
    else:
        base_font, bold_font = "Helvetica", "Helvetica-Bold"

    getSampleStyleSheet()
    title_style = ParagraphStyle("title", fontName=bold_font, fontSize=16, spaceAfter=6)
    meta_style = ParagraphStyle("meta", fontName=base_font, fontSize=10,
                                textColor=colors.HexColor("#666666"), spaceAfter=12)
    cell_style = ParagraphStyle("cell", fontName=base_font, fontSize=10)
    header_style = ParagraphStyle("h", fontName=bold_font, fontSize=10)

    table_data = [[Paragraph("Параметр", header_style), Paragraph("Значение", header_style)]]
    for key, value in get_rows(sat):
        table_data.append([Paragraph(key, cell_style), Paragraph(value, cell_style)])

    t = Table(table_data, colWidths=[50*mm, 120*mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F5F5F5")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))

    doc.build([
        Paragraph(sat.title_content or "", title_style),
        Paragraph(sat.small_content or "", meta_style),
        t,
    ])
    buf.seek(0)
    filename = f"{sat.title_content or 'object'}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": content_disposition(filename)},
    )


def export_html(sat: SSatelliteOut) -> StreamingResponse:
    rows_html = "\n".join(
        f"<tr><td>{escape(k)}</td><td>{escape(v)}</td></tr>"
        for k, v in get_rows(sat)
    )
    title = escape(sat.title_content or '')
    description = escape(sat.small_content or '')
    html = f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  body{{font-family:Arial,sans-serif;max-width:800px;margin:40px auto;color:#222}}
  h1{{margin-bottom:4px}}
  p{{color:#666;margin-top:0}}
  table{{border-collapse:collapse;width:100%;margin-top:16px}}
  th,td{{border:1px solid #ccc;padding:8px 12px;text-align:left}}
  th{{background:#f5f5f5;font-weight:600}}
</style>
</head>
<body>
<h1>{title}</h1>
<p>{description}</p>
<table>
<thead><tr><th>Параметр</th><th>Значение</th></tr></thead>
<tbody>{rows_html}</tbody>
</table>
</body>
</html>"""
    filename = f"{sat.title_content or 'object'}.html"
    return StreamingResponse(
        io.BytesIO(html.encode("utf-8")),
        media_type="text/html; charset=utf-8",
        headers={"Content-Disposition": content_disposition(filename)},
    )
